"""문서 파서: HWPX / DOCX / PDF / TXT·MD → DocumentBlock 목록.

- HWPX는 ZIP 안의 Contents/section*.xml(OWPML)을, DOCX는 본문 순서대로 문단·표를 읽는다.
- PDF는 pdfplumber로 표 영역과 본문을 분리한다. 페이지를 넘어 이어지는 표는 하나로 합친다.
  (pdfplumber가 없으면 pypdf로 글자만 읽는다)
- TXT·MD는 UTF-8 → CP949 순으로 인코딩을 시도한다.
- 표는 형식과 관계없이 행·셀 그대로 넘겨 같은 규칙으로 요구사항 표를 인식한다.
    · 상세형: 'ID 계열 칸(고유번호·요구사항 ID 등) | 코드 값(SFR-001, MAR-001…)'이 있는 키-값 표 → 표 1개 = 요구사항 1개
    · 목록형: 머리글에 ID 열과 세부내용 열이 있는 표 → 행 1개 = 요구사항 1개로 나눈다
- 제목 경로: Ⅰ/[붙임]/[별지] → 1. → 가. 계층을 블록마다 붙인다. 목차 줄(끝에 쪽번호)은 제목으로 보지 않는다.
- 위치 정보: HWPX·DOCX에는 페이지 정보가 없으므로 섹션·순번·제목 경로로 표시한다. PDF는 페이지 번호도 남긴다.
"""
from __future__ import annotations

import hashlib
import logging
import re
import zipfile
from functools import lru_cache
from pathlib import Path

from lxml import etree

from app.schemas.document import BlockLocation, DocumentBlock

log = logging.getLogger(__name__)

HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"

ROMAN = "ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ"
_L1 = re.compile(rf"^([{ROMAN}]|[IVX]+\.)\s*\S|^첨부\s*\d+")
# 붙임·별지·별표·서식 제목: '[붙임1] …', '[별지 제13호] …', '별표 2. …' → 장 단위로 본다 (서식 안의 '1. …'이 제목으로 새지 않게)
_ATTACH = re.compile(r"^\[(붙임|별지|별표|서식|첨부)[^\]]{0,20}\]|^(붙임|별지|별표)\s*(제\s*)?\d+(-\d+)?\s*호?[.)]?\s")
_L2 = re.compile(r"^\d{1,2}\.\s*\S")
_L3 = re.compile(r"^[가-하]\.\s*\S")
_SENTENCE_END = re.compile(r"(다|함|음|임)\.?$")    # '1. …해야 한다.' 같은 문장은 제목이 아님 ('개요'는 제목)
_TOC_ENTRY = re.compile(r"([가-힣A-Za-z)\]]\s?|\.{3,}\s*|·{3,}\s*|…+\s*)\d{1,3}$")  # 목차 줄: '사업 개요4', '시스템 현황 10'
_ITEM_MARKS = ("□", "■")                           # RFP의 최상위 항목 기호: 새 문단의 시작으로 본다
# 요구사항 코드 값: SFR-001, MAR-001, REQ_01, FR01, SFR-001-01 …
REQ_CODE = re.compile(r"^[A-Za-z]{2,6}[-_ ]?\d{1,4}([-_.]\d{1,3})?$")

SUPPORTED_SUFFIXES = {".hwpx", ".docx", ".pdf", ".txt", ".md"}
# 읽지 못하는 형식은 사용자가 바로 조치할 수 있게 변환 방법을 알려준다
CONVERT_HINTS = {
    ".hwp": "한글에서 열고 [다른 이름으로 저장] → HWPX 또는 PDF로 저장해서 넣으세요. "
            "구 .hwp는 한컴 독자 바이너리 형식이라 직접 읽지 않습니다.",
    ".hwt": "한글 서식 파일입니다. 내용이 들어 있는 .hwpx 문서를 넣으세요.",
    ".doc": "Word에서 열고 .docx 또는 PDF로 저장해서 넣으세요.",
    ".ppt": "PowerPoint에서 열고 PDF로 저장해서 넣으세요.",
    ".pptx": "PowerPoint에서 열고 PDF로 저장해서 넣으세요.",
    ".xls": "Excel에서 열고 PDF로 저장하거나, 표 내용을 텍스트로 옮겨 넣으세요.",
    ".xlsx": "Excel에서 열고 PDF로 저장하거나, 표 내용을 텍스트로 옮겨 넣으세요.",
    ".zip": "압축을 풀고 안의 문서를 한 건씩 넣으세요.",
}

# 요구사항 표의 머리 칸(공백·기호 제거, 소문자) → 표준 필드명. 발주기관마다 표기가 달라 동의어를 모아 둔다.
FIELD_KEYS = {
    "고유번호": "code", "요구사항고유번호": "code", "요구사항번호": "code", "요구사항id": "code", "id": "code",
    "요구사항식별번호": "code", "요구사항식별자": "code", "식별번호": "code", "요구사항코드": "code",
    "명칭": "name", "요구사항명칭": "name", "요구사항명": "name", "요구사항이름": "name",
    "분류": "category", "요구사항분류": "category", "요구사항유형": "category", "유형": "category", "구분": "category",
    "응락수준": "acceptance", "수용수준": "acceptance", "중요도": "acceptance", "우선순위": "acceptance",
    "정의": "definition", "요구사항정의": "definition",
    "세부내용": "details", "내용": "details", "상세내용": "details", "요구사항내용": "details", "요구사항세부내용": "details",
    "요구사항상세내용": "details", "상세설명": "details", "요구사항상세설명": "details",
    "산출정보": "outputs", "산출물": "outputs", "관련요구사항": "related", "출처": "origin", "요구사항출처": "origin",
}
FIELD_LABELS = {
    "code": "고유번호", "name": "명칭", "category": "분류", "acceptance": "응락 수준",
    "definition": "정의", "details": "세부내용", "outputs": "산출정보", "related": "관련 요구사항", "origin": "출처",
}


class UnsupportedFormatError(ValueError):
    """읽을 수 없는 파일 형식. hint에 사용자가 할 수 있는 변환 방법을 담는다."""

    def __init__(self, suffix: str, hint: str, message: str = ""):
        super().__init__(message or f"지원하지 않는 형식입니다: {suffix or '(확장자 없음)'}")
        self.suffix = suffix
        self.hint = hint


def check_supported(path: str | Path) -> None:
    """파싱 전에 형식을 확인한다. 지원하지 않으면 변환 안내와 함께 UnsupportedFormatError를 던진다.

    PDF는 글자 층도 확인한다: 스캔본(글자 없는 페이지가 절반 초과)은 OCR 없이 입력 단계에서 막는다.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        hint = CONVERT_HINTS.get(suffix, f"지원 형식: {', '.join(sorted(SUPPORTED_SUFFIXES))}")
        raise UnsupportedFormatError(suffix, hint)
    if suffix == ".pdf":
        _check_scanned_pdf(path)


def file_fingerprint(path: Path) -> tuple[str, str]:
    """파일 내용 해시로 (document_id, document_version)을 만든다. 같은 파일이면 항상 같은 값."""
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    return f"DOC-{digest[:10]}", digest[:16]


def read_text_file(path: Path) -> str:
    """텍스트 파일을 UTF-8(BOM 포함) → CP949 순으로 읽는다. 둘 다 실패하면 깨진 글자를 대체해 읽는다."""
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding != "utf-8-sig":
            log.info("인코딩을 %s로 읽었습니다: %s", encoding, path.name)
        return text
    log.warning("인코딩을 알 수 없어 깨진 글자를 대체해 읽었습니다: %s", path.name)
    return raw.decode("utf-8", errors="replace")


def parse_document(path: str | Path) -> list[DocumentBlock]:
    """파일 하나를 형식에 맞는 파서로 읽어, 원문 위치가 붙은 DocumentBlock 목록으로 돌려준다 (파서 진입점)."""
    path = Path(path)
    check_supported(path)
    suffix = path.suffix.lower()
    document_id, version = file_fingerprint(path)
    if suffix == ".hwpx":
        raw = _parse_hwpx(path)
    elif suffix == ".docx":
        raw = _parse_docx(path)
    elif suffix in {".txt", ".md"}:
        raw = _parse_text(read_text_file(path), markdown=suffix == ".md")
    else:
        raw = _parse_pdf(path)
    return _finalize(_repair_requirements(raw), document_id, version)


_CODE_MENTION = re.compile(r"(?<![A-Za-z])([A-Z]{2,6}-\d{1,4})(?!\d)")


def document_codes(path: str | Path) -> set[str] | None:
    """원문 전체(글상자·중첩 표 포함)에 나오는 요구사항 번호 모양 글자. 파싱에서 빠진 요구사항을 찾는 데 쓴다.

    HWPX·DOCX는 XML의 모든 문단을 직접 훑는다. PDF·TXT는 파서가 읽은 글자가 전부이므로 None.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    texts: list[str] = []
    if suffix == ".hwpx":
        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if re.fullmatch(r"Contents/section\d+\.xml", name):
                    root = etree.fromstring(zf.read(name))
                    texts += ["".join(t.text or "" for t in p.findall(f"{HP}run/{HP}t")) for p in root.iter(f"{HP}p")]
    elif suffix == ".docx":
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as zf:
            root = etree.fromstring(zf.read("word/document.xml"))
            texts += ["".join(t.text or "" for t in p.iter(f"{w}t")) for p in root.iter(f"{w}p")]
    else:
        return None
    return {m for text in texts for m in _CODE_MENTION.findall(text)}


def heading_level(text: str) -> int | None:
    """문단이 RFP식 제목이면 단계(0: Ⅰ·첨부·[붙임]·[별지], 1: 1., 2: 가.)를, 아니면 None을 돌려준다."""
    # 쉼표로 끝나면 줄바꿈된 본문 문장의 첫 줄이다 (PDF에서 '가. 국가기관, …'처럼 잘린 줄)
    if len(text) > 60 or _SENTENCE_END.search(text) or _TOC_ENTRY.search(text) or text.endswith(","):
        return None
    if _L1.match(text) or _ATTACH.match(text):
        return 0
    if _L2.match(text):
        return 1
    if _L3.match(text):
        return 2
    return None


# ------------------------------------------------------------------ HWPX
def _parse_hwpx(path: Path) -> list[dict]:
    """HWPX의 섹션 XML을 번호순으로 읽어 문단·표 원시 항목 목록을 만든다."""
    blocks: list[dict] = []
    with zipfile.ZipFile(path) as zf:
        sections = sorted(
            (n for n in zf.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)),
        )
        if not sections:
            raise ValueError("HWPX 본문(Contents/section*.xml)을 찾을 수 없습니다.")
        for sec_idx, name in enumerate(sections):
            root = etree.fromstring(zf.read(name))
            for p in root.findall(f"{HP}p"):
                blocks.extend(_walk_paragraph(p, sec_idx))
    return blocks


def _run_text(p) -> str:
    """HWPX 문단(hp:p) 하나의 글자를 이어 붙인다."""
    parts = []
    for run in p.findall(f"{HP}run"):
        for t in run.findall(f"{HP}t"):
            parts.append("".join(t.itertext()))
    return "".join(parts).strip()


def _walk_paragraph(p, section: int) -> list[dict]:
    """문단 하나에서 본문 글자, 문단에 붙은 표(hp:tbl), 글상자(도형 속 hp:drawText)를 차례로 원시 항목으로 뽑는다.

    글상자 안의 문단·표도 본문처럼 읽는다. PDF를 HWPX로 변환한 문서는 내용 일부(요구사항 표 포함)를 글상자로 만든다.
    """
    out = []
    text = _run_text(p)
    if text:
        out.append({"type": "paragraph", "section": section, "text": text})
    for run in p.findall(f"{HP}run"):
        for child in run:
            if child.tag == f"{HP}tbl":
                out.append({"type": "table", "section": section, "rows": _table_rows(child)})
            elif child.tag != f"{HP}t":
                for sub in _text_boxes(child):
                    for q in sub.findall(f"{HP}p"):
                        out.extend(_walk_paragraph(q, section))
    return out


def _text_boxes(shape) -> list:
    """도형(사각형·다각형·묶음 등) 안 글상자의 subList 목록. 글상자 속 글상자는 재귀 호출에서 읽으므로 바깥 것만."""
    found = []
    for dt in shape.iter(f"{HP}drawText"):
        parent = dt.getparent()
        while parent is not None and parent is not shape and parent.tag != f"{HP}drawText":
            parent = parent.getparent()
        if parent is not None and parent.tag == f"{HP}drawText":
            continue
        if (sub := dt.find(f"{HP}subList")) is not None:
            found.append(sub)
    return found


def _table_rows(tbl) -> list[list[str]]:
    """표를 행·셀 목록으로 읽는다. 병합된 셀은 실제 셀만 남는다."""
    rows = []
    for tr in tbl.findall(f"{HP}tr"):
        rows.append([_cell_text(tc) for tc in tr.findall(f"{HP}tc")])
    return rows


def _cell_text(tc) -> str:
    """셀 안의 문단을 줄바꿈으로 잇는다. 셀 안의 중첩 표는 'a | b' 한 줄씩으로 펼친다."""
    sub = tc.find(f"{HP}subList")
    if sub is None:
        return ""
    lines = []
    for p in sub.findall(f"{HP}p"):
        text = _run_text(p)
        if text:
            lines.append(text)
        for run in p.findall(f"{HP}run"):
            for nested in run.findall(f"{HP}tbl"):
                lines.extend(" | ".join(c for c in row if c) for row in _table_rows(nested))
    return "\n".join(lines).strip()


# ------------------------------------------------------------------ DOCX
def _parse_docx(path: Path) -> list[dict]:
    """DOCX 본문을 문단·표 순서 그대로 읽는다. Word 제목 스타일(Heading 1~3, 제목 1~3)은 제목 단계로 넘긴다."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[dict] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            item = {"type": "paragraph", "section": 0, "text": text}
            level = _docx_heading_level(para.style.name if para.style is not None else "")
            if level is not None:
                item["heading_level"] = level
            blocks.append(item)
        elif tag == "tbl":
            blocks.append({"type": "table", "section": 0, "rows": _docx_rows(Table(child, doc))})
    return blocks


def _docx_heading_level(style_name: str) -> int | None:
    """Word 스타일 이름에서 제목 단계를 읽는다 ('Heading 2'·'제목 2' → 1)."""
    m = re.match(r"(heading|제목)\s*(\d)", style_name or "", re.I)
    return min(int(m.group(2)), 3) - 1 if m else None


def _docx_rows(table) -> list[list[str]]:
    """DOCX 표를 행·셀 목록으로 읽는다. 가로 병합으로 반복되는 같은 셀은 한 번만 넣는다 (HWPX와 같은 모양)."""
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(_docx_cell_text(cell))
        rows.append(cells)
    return rows


def _docx_cell_text(cell) -> str:
    """셀 안 문단을 줄바꿈으로 잇고, 중첩 표는 'a | b' 한 줄씩으로 펼친다."""
    lines = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    for nested in cell.tables:
        lines.extend(" | ".join(c for c in row if c) for row in _docx_rows(nested))
    return "\n".join(lines).strip()


# ------------------------------------------------------------------ TXT / MD
def _parse_text(text: str, markdown: bool = False, page: int | None = None) -> list[dict]:
    """줄 단위로 읽어 제목 줄은 단독 블록으로, 나머지는 빈 줄·'□' 항목을 경계로 문단 블록으로 묶는다.

    제목 줄을 따로 떼어야 PDF처럼 빈 줄이 없는 본문에서도 제목 경로(heading_path)가 잡힌다.
    """
    blocks: list[dict] = []
    buffer: list[str] = []

    def flush():
        if buffer:
            blocks.append({"type": "paragraph", "section": 0, "text": "\n".join(buffer), "page": page})
            buffer.clear()

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            flush()
            continue
        md = re.match(r"^(#{1,6})\s+(.+)", line) if markdown else None
        if md:
            flush()
            blocks.append({"type": "paragraph", "section": 0, "text": md.group(2).strip(), "page": page,
                           "heading_level": min(len(md.group(1)), 3) - 1})
            continue
        if heading_level(line) is not None:
            flush()
            blocks.append({"type": "paragraph", "section": 0, "text": line, "page": page})
            continue
        if line.startswith(_ITEM_MARKS):
            flush()
        buffer.append(line)
    flush()
    return blocks


# ------------------------------------------------------------------ PDF
def _parse_pdf(path: Path) -> list[dict]:
    """PDF를 페이지별로 읽는다. pdfplumber가 있으면 표와 본문을 분리하고, 없으면 글자만 읽는다.

    스캔본(글자 층이 없는 PDF)은 check_supported에서 미리 막는다. 글자가 없는 일부 페이지(그림 등)는 건너뛴다.
    """
    try:
        import pdfplumber
    except ImportError:
        log.warning("pdfplumber가 없어 PDF 표 구조 없이 글자만 읽습니다: pip install pdfplumber")
        return _parse_pdf_text_only(path)

    blocks: list[dict] = []
    empty_pages: list[int] = []
    with pdfplumber.open(str(path)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            items = _pdf_page_items(page)
            if not items:
                empty_pages.append(page_no)
                continue
            for i, (kind, value) in enumerate(items):
                if kind == "table":
                    rows = value
                    # 페이지 맨 앞의 표가 직전 페이지 마지막 표에 이어지면(머리글 없이 시작) 하나로 합친다
                    if i == 0 and blocks and blocks[-1]["type"] == "table" \
                            and blocks[-1].get("page_end") == page_no - 1 and not _is_requirement_table(rows):
                        blocks[-1]["rows"].extend(rows)
                        blocks[-1]["page_end"] = page_no
                        continue
                    blocks.append({"type": "table", "section": 0, "rows": rows, "page": page_no,
                                   "page_end": page_no})
                elif kind == "title":
                    blocks.append({"type": "paragraph", "section": 0, "text": value, "page": page_no,
                                   "heading_level": 0})
                else:
                    blocks.extend(_parse_text(value, page=page_no))
    if empty_pages:
        log.warning("글자가 없는 페이지는 건너뜀 (그림·스캔 페이지): %s p.%s", path.name,
                    ",".join(map(str, empty_pages)))
    return blocks


_PAGE_NUMBER = re.compile(r"^[-–—(]?\s*\d{1,4}\s*[-–—)]?$|^\d{1,4}\s*/\s*\d{1,4}$")   # '- 11 -', '(3)', '3 / 42'
# 쪽 가운데에 있어도 쪽 번호로 보는 꼴 ('- 18 -'만). 인쇄한 PDF(Microsoft Print to PDF 등)는 원래 문서의 쪽 경계가
# PDF 쪽 가운데에 와서, 원래 꼬리말 쪽 번호가 잘린 요구사항 표의 앞뒤 사이에 낀다. '(3)'·'3'은 본문 번호일 수 있어 여백에서만.
_PAGE_NUMBER_DASHED = re.compile(r"^[-–—]\s*\d{1,4}\s*[-–—]$")
_LATIN_ROMAN = {"I": "Ⅰ", "II": "Ⅱ", "III": "Ⅲ", "IV": "Ⅳ", "V": "Ⅴ", "VI": "Ⅵ", "VII": "Ⅶ", "VIII": "Ⅷ",
                "IX": "Ⅸ", "X": "Ⅹ"}


def _contains(outer: tuple, inner: tuple, tol: float = 1.0) -> bool:
    return (outer[0] - tol <= inner[0] and inner[2] <= outer[2] + tol
            and outer[1] - tol <= inner[1] and inner[3] <= outer[3] + tol)


def _title_box_text(rows: list[list[str]]) -> str | None:
    """'Ⅱ | 제 | 안 | 요 | 구 | 사 | 항'처럼 글자를 한 칸씩 넣은 장 제목 상자면 'Ⅱ 제안요구사항'을 돌려준다.

    표로 잡히지 않으면 같은 제목이 'Ⅱ 제 안 요 구 사 항' 한 줄로 나오므로 그 줄도 rows=[[줄]]로 받아 처리한다.
    """
    cells = [c for r in rows for c in r if c]
    if len(cells) == 1:
        cells = cells[0].split()
    if not 3 <= len(cells) <= 30 or any(len(c) > 1 for c in cells[1:]) or len(cells[0]) > 4:
        return None
    numeral = _LATIN_ROMAN.get(cells[0], cells[0])
    if numeral not in ROMAN:
        return None
    return f"{numeral} {''.join(cells[1:])}"


def _open_side_edges(page, tol: float = 3.0) -> list[dict]:
    """좌우 바깥 세로선이 없는 표(가로선만 끝까지 긋는 공문서 양식)의 양 끝 세로선을 보충한다.

    표 안쪽 세로선이 차지하는 높이를 표 하나의 범위로 보고, 그 범위의 가로선 양 끝에 세로선이 없으면 더한다.
    바깥 세로선이 없으면 pdfplumber가 값 칸을 닫지 못해 '요구사항 명칭 | (빈칸)'처럼 값이 빠진다.
    """
    h = [e for e in page.edges if e["orientation"] == "h" and e["x1"] - e["x0"] > 20]
    v = [e for e in page.edges if e["orientation"] == "v" and e["bottom"] - e["top"] > 5]
    if not h or not v:
        return []
    left, right = min(e["x0"] for e in h), max(e["x1"] for e in h)
    bands: list[list[float]] = []
    for e in sorted((e for e in v if left + tol < e["x0"] < right - tol), key=lambda e: e["top"]):
        if bands and e["top"] <= bands[-1][1] + 2:
            bands[-1][1] = max(bands[-1][1], e["bottom"])
        else:
            bands.append([e["top"], e["bottom"]])
    extra: list[dict] = []
    for top, bottom in bands:
        rows = [e for e in h if top - 2 <= e["top"] <= bottom + 2]
        if len(rows) < 2:
            continue
        for x in (min(e["x0"] for e in rows), max(e["x1"] for e in rows)):
            covered = sum(max(0.0, min(e["bottom"], bottom) - max(e["top"], top))
                          for e in v if abs(e["x0"] - x) <= tol)
            if covered < (bottom - top) * 0.8:
                extra.append({"object_type": "line", "x0": x, "x1": x, "top": top, "bottom": bottom,
                              "width": 0, "height": bottom - top})
    return extra


def _pdf_page_items(page) -> list[tuple[str, object]]:
    """한 페이지를 위→아래 순서의 ('text', 본문) / ('table', 행 목록) / ('title', 장 제목) 항목으로 나눈다.

    - 그림자 효과로 1~2pt 어긋나게 두 번 찍힌 글자('사사업업')는 하나로 합친다.
    - 다른 표 안에 들어 있는 표(셀 속 중첩 표)는 바깥 표 셀 글자에 이미 들어 있으므로 따로 꺼내지 않는다.
    - 글자를 한 칸씩 넣은 장 제목 상자는 표가 아니라 최상위 제목으로 돌려준다.
    - 머리말·꼬리말 영역의 쪽번호('- 11 -')는 버린다.
    - 좌우 바깥 세로선이 없는 표는 양 끝 세로선을 보충해서 찾는다.
    """
    page = page.dedupe_chars(tolerance=2)
    extra = _open_side_edges(page)
    found = page.find_tables({"explicit_vertical_lines": extra}) if extra else page.find_tables()
    nested = [t for t in found if any(o is not t and _contains(o.bbox, t.bbox) for o in found)]
    tables = sorted((t for t in found if t not in nested), key=lambda t: t.bbox[1])
    boxes = [t.bbox for t in tables]

    def outside(obj) -> bool:
        # 글자 중심으로 판단한다 (pdfplumber가 셀에 글자를 넣는 기준과 같게). 표 아래 테두리에 걸친 마지막 줄이
        # 셀에도 들어가고 본문에도 한 번 더 나오는 것을 막는다.
        x0, top, x1, bottom = obj.get("x0", 0), obj.get("top", 0), obj.get("x1", 0), obj.get("bottom", 0)
        return not any(_center_in({"x0": x0, "top": top, "x1": x1, "bottom": bottom}, b) for b in boxes)

    def page_number(line: dict) -> bool:
        if _PAGE_NUMBER_DASHED.match(line["text"].strip()):
            return True
        margin = page.height * 0.08
        in_margin = line["top"] > page.height - margin or line["bottom"] < margin
        return in_margin and bool(_PAGE_NUMBER.match(line["text"].strip()))

    body = page.filter(outside) if boxes else page
    positioned: list[tuple[float, str, object]] = []
    for line in body.extract_text_lines():
        if not line["text"].strip() or page_number(line):
            continue
        title = _title_box_text([[line["text"].strip()]])
        positioned.append((line["top"], "title", title) if title else (line["top"], "line", line["text"]))
    for table in tables:
        rows = _clean_pdf_rows(_extract_with_nested(page, table, [t for t in nested if _contains(table.bbox, t.bbox)]))
        if not rows:
            continue
        if title := _title_box_text(rows):
            positioned.append((table.bbox[1], "title", title))
        else:
            positioned.append((table.bbox[1], "table", rows))
    positioned.sort(key=lambda x: x[0])

    items: list[tuple[str, object]] = []
    for _, kind, value in positioned:
        if kind == "line":
            if items and items[-1][0] == "text":
                items[-1] = ("text", f"{items[-1][1]}\n{value}")
            else:
                items.append(("text", value))
        else:
            items.append((kind, value))
    return items


def _extract_with_nested(page, table, inner: list) -> list[list[str | None]]:
    """표를 행 목록으로 꺼낸다. 셀 속 중첩 표는 그 자리에 '칸 | 칸' 줄로 바꿔 넣는다.

    pdfplumber는 셀 글자를 줄 단위로 읽어 중첩 표의 칸이 서로 섞인다
    ('기본 회원정보 이름 / 아이디 …' 다음 줄에 '(필수)'). 그러면 항목명과 값의 대응이 깨진다.
    """
    rows = table.extract()
    if not inner:
        return rows

    def text_in(x0: float, top: float, x1: float, bottom: float) -> str:
        if bottom <= top:
            return ""
        return page.filter(lambda o: o.get("object_type") == "char"
                           and _center_in(o, (x0, top, x1, bottom))).extract_text() or ""

    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            subs = sorted((t for t in inner if cell and _contains(cell, t.bbox)), key=lambda t: t.bbox[1])
            if not subs:
                continue
            parts, top = [], cell[1]
            for sub in subs:
                parts.append(text_in(cell[0], top, cell[2], sub.bbox[1]))
                parts += [" | ".join(" ".join((x or "").split()) for x in sub_row if x is not None)
                          for sub_row in sub.extract()]
                top = sub.bbox[3]
            parts.append(text_in(cell[0], top, cell[2], cell[3]))
            rows[r][c] = "\n".join(p for p in parts if p.strip())
    return rows


def _center_in(obj: dict, bbox: tuple) -> bool:
    cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
    return bbox[0] <= cx < bbox[2] and bbox[1] <= cy < bbox[3]


# ------------------------------------------------------------------ 스캔본 판별
TEXTLESS_MIN_CHARS = 20     # 글자가 이보다 적은 페이지는 '글자 없는 페이지'(그림·스캔)로 본다
SCANNED_MAX_RATIO = 0.5     # 글자 없는 페이지가 이 비율을 넘으면 스캔본으로 보고 입력을 막는다
SCANNED_HINT = ("스캔본이나 이미지로 저장된 PDF는 지원하지 않습니다. 원본 프로그램(한글·Word 등)에서 "
                "'PDF로 저장/내보내기'로 다시 만들어 주세요. ('Microsoft Print to PDF'로 인쇄하면 글자가 그림이 될 수 있습니다)")


@lru_cache(maxsize=32)
def _textless_pages_cached(path: str, mtime_ns: int, size: int) -> tuple[tuple[int, ...], int]:
    """(글자 없는 페이지 번호들, 전체 페이지 수). 판별만 하므로 pdfplumber보다 2배쯤 빠른 pypdf를 쓴다.
    mtime_ns·size는 파일이 바뀌면 캐시를 새로 쓰게 하는 키."""
    from pypdf import PdfReader

    pages = PdfReader(path).pages
    empty = [i for i, p in enumerate(pages, 1) if len((p.extract_text() or "").strip()) < TEXTLESS_MIN_CHARS]
    return tuple(empty), len(pages)


def textless_pages(path: str | Path) -> list[int]:
    """PDF에서 글자가 (거의) 없는 페이지 번호 목록. PDF가 아니면 빈 목록."""
    path = Path(path)
    if path.suffix.lower() != ".pdf" or not path.exists():
        return []
    stat = path.stat()
    return list(_textless_pages_cached(str(path), stat.st_mtime_ns, stat.st_size)[0])


def _check_scanned_pdf(path: Path) -> None:
    """글자 없는 페이지가 절반을 넘으면 스캔본으로 보고 UnsupportedFormatError를 던진다."""
    if not path.exists():
        return
    try:
        stat = path.stat()
        empty, total = _textless_pages_cached(str(path), stat.st_mtime_ns, stat.st_size)
    except Exception as e:  # noqa: BLE001 — 깨진 PDF는 파서에서 오류를 내게 둔다
        log.debug("PDF 글자 확인 실패: %s", e)
        return
    if total and len(empty) / total > SCANNED_MAX_RATIO:
        raise UnsupportedFormatError(
            ".pdf", f"{SCANNED_HINT} (전체 {total}쪽 중 {len(empty)}쪽에 글자가 없음)",
            message=f"스캔본 PDF는 지원하지 않습니다: {path.name}")


def _clean_pdf_rows(rows: list[list[str | None]]) -> list[list[str]]:
    """pdfplumber 표 행을 정리한다. None(병합으로 가려진 칸)은 빼서 HWPX처럼 실제 셀만 남긴다."""
    cleaned = []
    for row in rows:
        cells = [(c or "").strip() for c in row if c is not None]
        if any(cells):
            cleaned.append(cells)
    return cleaned


def _parse_pdf_text_only(path: Path) -> list[dict]:
    """pdfplumber가 없을 때의 대체 경로: pypdf로 페이지 글자만 읽는다 (표 구조 없음)."""
    from pypdf import PdfReader

    blocks = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        blocks.extend(_parse_text(page.extract_text() or "", page=i))
    return blocks


# ------------------------------------------------------------------ 공통 후처리
def _norm_key(key: str) -> str:
    """표 머리 칸 비교용: 공백·/·:·가운뎃점을 지우고 소문자로 ('세부 / 내용' → '세부내용', '요구사항 ID' → '요구사항id')."""
    return re.sub(r"[\s/·:]", "", key).lower()


_SINGLE_FIELDS = {"code", "name", "category", "acceptance"}
_TEXT_FIELDS = {"definition", "details", "outputs", "related", "origin"}


def _trim(row: list[str]) -> list[str]:
    """행 앞뒤의 빈 칸을 지운다 (PDF→HWPX 변환본의 ['', '요구사항 번호', 'SFR-02', …, ''] 같은 행)."""
    cells = [c.strip() for c in row]
    while cells and not cells[0]:
        cells.pop(0)
    while cells and not cells[-1]:
        cells.pop()
    return cells


def parse_requirement_table(rows: list[list[str]]) -> dict[str, str]:
    """'고유번호 | SFR-001', '분류 | 기능 | 응락 수준 | 필수', '설명 | 정의 | …' 형태의 표를 필드로 변환.

    - 칸 이름(FIELD_KEYS) 뒤의 칸들이 그 값이다. 칸 이름 앞의 상위 분류 칸('설명')은 무시한다.
    - 칸 이름이 없는 행(세부내용 안의 하위 표, 쪽 넘김으로 잘린 뒷부분)은 직전 서술 필드(기본: 세부내용)에 붙인다.
    - 세부내용이 나온 뒤 다시 나온 '구분 | 분류 | 중분류'는 하위 표 머리글이므로 분류를 덮지 않고 내용으로 본다.
    """
    fields: dict[str, str] = {}
    current = "details"

    def put(field: str, value: str):
        nonlocal current
        if value.strip():
            fields[field] = (fields.get(field, "") + "\n" + value.strip()).strip()
            if field in _TEXT_FIELDS:
                current = field

    for row in rows:
        cells = _trim(row)
        if not cells:
            continue
        keys = [FIELD_KEYS.get(_norm_key(c)) for c in cells]
        if "details" in fields or "definition" in fields:
            keys = [None if k in _SINGLE_FIELDS and k in fields else k for k in keys]
        positions = [i for i, k in enumerate(keys) if k and i < len(cells) - 1]
        if not positions:
            if not any(keys):
                put(current, " | ".join(cells))
            continue
        for n, i in enumerate(positions):
            end = positions[n + 1] if n + 1 < len(positions) else len(cells)
            put(keys[i], "\n".join(c for c in cells[i + 1:end] if c))
    return fields


# ------------------------------------------------------------------ 깨진 요구사항 표 복구
# PDF를 HWPX로 변환한 문서 등에서 보이는 모양 (정상 문서의 표는 손대지 않는다):
#   - 한 표에 요구사항 여러 개 (번호 행이 두 번 나옴)
#   - 칸 이름과 값이 한 셀에 붙음: 'SFR-03요구사항 분류시스템 기능 요구사항 데이터 금고 …'
#   - 표가 문단으로 풀림: '요구사항 번호 요구사항 이름 요구사항 내용' / 'SER-01요구사항 분류보안 요구사항 일반 요구사항' / 본문 …
#   - 쪽 넘김으로 잘린 뒷부분이 따로 표가 됨: ['', '', '…이어지는 내용', ''], ['산출정보', '…']
#     (인쇄한 PDF는 행 사이에서 잘려 칸 이름으로 시작하기도 함: ['세부내용', '…'], ['산출정보', '…'])
_LABEL = re.compile(r"요구사항\s*(?:고유\s*)?번호|요구사항\s*(?:ID|id)|요구사항\s*분류|요구사항\s*(?:이름|명칭)"
                    r"|요구사항\s*(?:상세\s*)?내용|세부\s*내용|산출\s*정보")
_GLUED = re.compile(r"(?<![A-Za-z])([A-Za-z]{2,6}-\d{1,4})\s*요구사항\s*분류")
_SUBHEADING = re.compile(r"^\(?\d{1,2}\)\s*\S|^\S.{0,40}(요구사항|제약사항)(\s*\([A-Za-z ]+\))?$")
_SYNTHETIC_LABELS = {"code": "요구사항 번호", "category": "요구사항 분류", "name": "요구사항 이름",
                     "details": "요구사항 내용", "outputs": "산출정보"}


def _code_row(cells: list[str]) -> bool:
    return any(FIELD_KEYS.get(_norm_key(c)) == "code" and REQ_CODE.match(cells[i + 1])
               for i, c in enumerate(cells[:-1]))


def _glued_cell(cells: list[str]) -> int | None:
    return next((i for i, c in enumerate(cells) if _GLUED.search(c)), None)


def _label_only(text: str) -> bool:
    return bool(text.strip()) and not _LABEL.sub("", text).strip()


_GLUED_LABEL = re.compile(rf"(?<!\s)({_LABEL.pattern})")   # 줄 맨 앞이거나 앞 글자에 붙은 칸 이름만 (문장 속 낱말은 제외)


def _fields_from_text(text: str) -> tuple[str, dict[str, str]]:
    """칸 이름과 값이 뒤섞인 글을 요구사항 필드로 읽는다. 돌려주는 값: (요구사항 앞의 글 — 예: 붙어 버린 절 제목, 필드)

    처리하는 모양
    - 붙은 셀:   '요구사항 번호PMR-01요구사항 분류프로젝트 관리 요구사항⏎요구사항 이름사업 관리⏎…'
    - 풀린 문단: '요구사항 번호 요구사항 이름 요구사항 내용' / 'SER-01요구사항 분류보안 요구사항 일반 요구사항' / 본문 …
    - 칸 이름과 값이 따로 온 글상자: '요구사항 번호' / '요구사항 이름' / 'PMR-05' / '운영자 매뉴얼' / …
    값이 없는 번호·이름 칸은 기다렸다가 다음 값에 채우고, 서술 칸('요구사항 내용', '산출정보') 뒤의 줄은 그 칸에 붙인다.
    이름 칸이 없으면 분류 값을 '…요구사항' 뒤에서 잘라 이름으로 쓴다 ('보안 요구사항 일반 요구사항' → 분류·이름).
    """
    starts = [m.start() for m in (_GLUED_LABEL.search(text), _GLUED.search(text)) if m]
    cut = text.rfind("\n", 0, min(starts)) + 1 if starts else 0
    before, body = text[:cut].strip(), text[cut:]
    fields: dict[str, str] = {}
    pending: list[str] = []
    current = "details"

    def assign(field: str, value: str):
        value = value.strip()
        if not value:
            return
        if field in _SINGLE_FIELDS and field not in fields:
            cm = re.match(r"(.+?(?:요구사항|사항))\s+(.+)$", value) if field == "category" else None
            if cm and "name" not in fields:
                fields["category"], fields["name"] = cm[1], cm[2]
                if "name" in pending:
                    pending.remove("name")
            else:
                fields[field] = value
            return
        target = current if field in _SINGLE_FIELDS else field
        fields[target] = (fields.get(target, "") + "\n" + value).strip()

    for line in body.split("\n"):
        line = line.strip()
        if _label_only(line):
            for label in _LABEL.findall(line):
                field = FIELD_KEYS.get(_norm_key(label), "details")
                if field in _TEXT_FIELDS:
                    current = field
                    fields.setdefault(field, "")     # 값이 없어도 칸이 있었음을 남긴다 (잘린 표 판별용)
                elif field not in fields and field not in pending:
                    pending.append(field)
            continue
        tokens = _GLUED_LABEL.split(line)
        if tokens[0].strip():
            if pending:
                assign(pending.pop(0), tokens[0])
            elif "code" not in fields and REQ_CODE.match(tokens[0].strip()):
                assign("code", tokens[0])
            else:
                assign(current, tokens[0])
        for label, value in zip(tokens[1::2], tokens[2::2]):
            field = FIELD_KEYS.get(_norm_key(label), "details")
            if field in _TEXT_FIELDS:
                current = field
                fields.setdefault(field, "")
            if value.strip():
                assign(field, value)
            elif field not in _TEXT_FIELDS and field not in fields:
                pending.append(field)
    return before, fields


def _synthetic_rows(fields: dict[str, str]) -> list[list[str]]:
    """복구한 필드를 정상 표 모양('칸 이름 | 값')으로. 산출정보는 값이 비어도 칸을 남긴다."""
    return [[_SYNTHETIC_LABELS[k], fields[k]] for k in _SYNTHETIC_LABELS
            if fields.get(k) or (k == "outputs" and k in fields)]


def _split_requirement_table(rows: list[list[str]]) -> list[tuple[str, object]] | None:
    """요구사항 표를 요구사항 하나씩으로 나눈다. 고칠 것이 없으면 None.

    번호 행('요구사항 번호 | SER-03')이나 붙은 셀('SFR-03요구사항 분류…')에서 새 요구사항이 시작된다.
    돌려주는 항목: ('req', 행 목록) / ('table', 요구사항 앞의 행) / ('text', 셀에 붙어 있던 절 제목)
    """
    trimmed = [_trim(r) for r in rows]
    starts = [i for i, cells in enumerate(trimmed) if _code_row(cells) or _glued_cell(cells) is not None]
    if not starts or (len(starts) == 1 and starts[0] <= 1 and _code_row(trimmed[starts[0]])):
        return None
    parts: list[tuple[str, object]] = []
    head = [r for r in trimmed[:starts[0]] if r]
    if head:
        parts.append(("table", head))
    for s, e in zip(starts, starts[1:] + [len(trimmed)]):
        seg = [r for r in trimmed[s:e] if r]
        cells = seg[0]
        g = None if _code_row(cells) else _glued_cell(cells)
        if g is None:
            parts.append(("req", seg))
            continue
        before, fields = _fields_from_text(cells[g])
        before = "\n".join([c for c in cells[:g] if not _label_only(c)] + ([before] if before else []))
        if before:
            parts.append(("text", before))
        extra = [c for c in cells[g + 1:] if not _label_only(c)]
        parts.append(("req", _synthetic_rows(fields) + ([extra] if extra else []) + seg[1:]))
    return parts


def _is_continuation(rows: list[list[str]]) -> bool:
    """쪽 넘김으로 잘린 요구사항 표의 뒷부분: 첫 칸이 비어 있고, 번호 행은 없고, '산출정보'·'세부내용' 칸이 있다."""
    if not rows or not rows[0] or rows[0][0].strip():
        return False
    trimmed = [_trim(r) for r in rows]
    if any(_code_row(c) or _glued_cell(c) is not None for c in trimmed):
        return False
    return any(c and FIELD_KEYS.get(_norm_key(c[0])) in ("outputs", "details") for c in trimmed)


def _continues_with_labels(rows: list[list[str]], head: list[list[str]]) -> bool:
    """쪽 넘김이 행 사이에 걸려 뒷부분이 칸 이름으로 시작하는 경우: ['세부내용', …], ['산출정보', …] (인쇄한 PDF 등).

    앞부분(head)이 산출정보 없이 끝났고, 뒷부분 행의 첫 칸이 모두 요구사항 칸 이름이며(세로 병합으로 빈 칸은 허용),
    번호·명칭처럼 하나뿐인 칸이 앞부분과 겹치지 않을 때만 이어진 것으로 본다 ('구분 | 내용' 같은 일반 표는 붙이지 않음).
    """
    if _has_outputs_row(head) or not rows or not rows[0] or not rows[0][0].strip():
        return False
    labels = [FIELD_KEYS.get(_norm_key(r[0])) for r in rows if r and r[0].strip()]
    if None in labels or "code" in labels:
        return False
    present = {FIELD_KEYS.get(_norm_key(c[0])) for c in map(_trim, head) if c}
    return not set(labels) & present & _SINGLE_FIELDS


def _repair_requirements(raw: list[dict]) -> list[dict]:
    """깨진 요구사항 표를 '요구사항 1개 = 표 1개'로 되돌린다 (정상 문서는 그대로 통과)."""
    def is_req(item: dict) -> bool:
        return item["type"] == "table" and _is_requirement_table(item["rows"])

    def paragraph(src: dict, text: str) -> dict:
        return {"type": "paragraph", "section": src.get("section", 0), "text": text, "page": src.get("page")}

    def cut_off_requirement() -> int | None:
        """잘린 표의 뒷부분을 붙일 요구사항: 바로 앞 요구사항 3개 중 '산출정보' 칸이 없는(=잘린) 가장 가까운 것.
        글상자 때문에 다른 요구사항이 사이에 끼는 경우가 있어 바로 앞 것만 보지 않는다."""
        candidates = []
        for k in range(len(out) - 1, -1, -1):
            if not is_req(out[k]) or len(candidates) == 3:
                break
            candidates.append(k)
        return next((k for k in candidates if not _has_outputs_row(out[k]["rows"])),
                    candidates[0] if candidates else None)

    out: list[dict] = []
    i = 0
    while i < len(raw):
        item = raw[i]
        i += 1
        if item["type"] == "table":
            if not any(c.strip() for r in item["rows"] for c in r):
                continue                        # 빈 표 (변환본의 틀만 남은 표)
            k = cut_off_requirement()
            if k is not None and (_is_continuation(item["rows"]) or _continues_with_labels(item["rows"], out[k]["rows"])):
                out[k] = {**out[k], "rows": out[k]["rows"] + [r for r in map(_trim, item["rows"]) if r]}
                continue
            parts = _split_requirement_table(item["rows"])
            if parts is None:
                out.append(item)
                continue
            for kind, value in parts:
                out.append(paragraph(item, value) if kind == "text" else {**item, "rows": value})
        elif _starts_requirement(raw, i - 1):
            # 문단으로 풀린 요구사항 표: 다음 표·제목·요구사항 전까지의 문단을 모은다
            start = i - 1
            has_code = _has_code_line(item["text"])
            while i < len(raw) and raw[i]["type"] == "paragraph" \
                    and not (has_code and _starts_requirement(raw, i)) \
                    and raw[i].get("heading_level", heading_level(raw[i]["text"])) is None \
                    and not (_SUBHEADING.match(raw[i]["text"]) and not _LABEL.search(raw[i]["text"])):
                has_code = has_code or _has_code_line(raw[i]["text"])
                i += 1
            before, fields = _fields_from_text("\n".join(r["text"] for r in raw[start:i]))
            if not fields.get("code"):          # 번호를 못 찾으면 원래 문단 그대로 둔다
                out.extend(raw[start:i])
                continue
            if before:
                out.append(paragraph(item, before))
            out.append({**paragraph(item, ""), "type": "table", "rows": _synthetic_rows(fields)})
        else:
            out.append(item)
    return out


def _has_outputs_row(rows: list[list[str]]) -> bool:
    return any(c and FIELD_KEYS.get(_norm_key(c[0])) == "outputs" for c in map(_trim, rows))


def _has_code_line(text: str) -> bool:
    return bool(_GLUED.search(text) or REQ_CODE.fullmatch(text.strip()))


def _starts_requirement(raw: list[dict], j: int) -> bool:
    """raw[j] 문단에서 풀린 요구사항이 시작되는지: 'SER-01요구사항 분류…' 같은 붙은 줄, 또는 칸 이름만 있는 줄
    ('요구사항 번호' / '요구사항 번호 요구사항 이름 요구사항 내용') 바로 뒤 3문단 안에 번호만 있는 줄·붙은 줄이 나올 때."""
    def para(k: int) -> str | None:
        return raw[k]["text"].strip() if k < len(raw) and raw[k]["type"] == "paragraph" else None

    text = para(j)
    if text is None:
        return False
    if _GLUED.search(text):
        return True
    if not (_label_only(text) and FIELD_KEYS.get(_norm_key(_LABEL.findall(text)[0])) == "code"):
        return False
    following = [t for t in (para(k) for k in range(j + 1, j + 4)) if t is not None]
    return any(_GLUED.search(t) or REQ_CODE.fullmatch(t) for t in following)


def render_requirement_table(fields: dict[str, str]) -> str:
    """요구사항 표 필드를 LLM과 사람이 읽기 좋은 텍스트로 만든다 (근거 인용문도 이 텍스트에서 확인)."""
    head = " | ".join(f"{FIELD_LABELS[k]}: {fields[k]}" for k in ("code", "name", "category", "acceptance") if k in fields)
    body = [f"[요구사항 표] {head}"]
    for k in ("definition", "details", "outputs", "related", "origin"):
        if k in fields:
            body.append(f"{FIELD_LABELS[k]}:\n{fields[k]}" if "\n" in fields[k] else f"{FIELD_LABELS[k]}: {fields[k]}")
    return "\n".join(body)


def render_table(rows: list[list[str]]) -> str:
    """일반 표를 'a | b | c' 줄 목록으로 만든다."""
    return "\n".join(" | ".join(c.replace("\n", " / ") for c in row) for row in rows if any(row))


def _is_requirement_table(rows: list[list[str]]) -> bool:
    """상세형 요구사항 표인지 판별: 앞 두 행 안에 'ID 계열 칸 | 코드 값' 쌍이 있으면 요구사항 표로 본다.

    예) '고유번호 | SFR-001' (ITSM 형식), '요구사항 ID | MAR-001 | 요구사항 명칭 | …' (KOBACO 형식)
    값이 코드 형태여야 하므로 '요구사항ID | 요구사항명칭 | …' 같은 빈 서식 머리글은 제외된다.
    """
    for row in rows[:2]:
        cells = [c.strip() for c in row]
        for i, cell in enumerate(cells[:-1]):
            if FIELD_KEYS.get(_norm_key(cell)) == "code" and REQ_CODE.match(cells[i + 1]):
                return True
    return False


def _align_rows(rows: list[list[str]], width: int) -> list[list[str]]:
    """세로 병합으로 왼쪽 칸이 빠진 행을 위 행 값으로 채워 머리글 열 수에 맞춘다."""
    aligned: list[list[str]] = []
    for row in rows:
        if len(row) < width and aligned:
            row = aligned[-1][:width - len(row)] + row
        aligned.append((row + [""] * width)[:width])
    return aligned


def _requirement_list_columns(rows: list[list[str]]) -> dict[int, str] | None:
    """목록형 요구사항 표(행 1개 = 요구사항 1개)면 {열 번호: 필드명}을, 아니면 None을 돌려준다.

    조건: 머리글에 ID 열과 세부내용(또는 정의) 열이 있고, 코드가 든 행이 2개 이상이며 세부내용이 평균 20자 이상.
    (ID·명칭만 나열한 요구사항 목록표는 상세 표와 대조하는 용도로 그대로 둔다)
    """
    if len(rows) < 3:
        return None
    mapping = {i: FIELD_KEYS[_norm_key(c)] for i, c in enumerate(rows[0]) if _norm_key(c) in FIELD_KEYS}
    fields = list(mapping.values())
    if "code" not in fields or not {"details", "definition"} & set(fields):
        return None
    code_col = next(i for i, f in mapping.items() if f == "code")
    detail_cols = [i for i, f in mapping.items() if f in ("details", "definition")]
    data = [r for r in _align_rows(rows[1:], len(rows[0])) if REQ_CODE.match(r[code_col].strip())]
    if len(data) < 2:
        return None
    avg = sum(len(r[c]) for r in data for c in detail_cols) / len(data)
    return mapping if avg >= 20 else None


def _finalize(raw: list[dict], document_id: str, version: str) -> list[DocumentBlock]:
    """원시 항목을 DocumentBlock으로 바꾼다. 블록 ID(B0001…)를 매기고, 제목을 따라가며 제목 경로를 붙인다."""
    blocks: list[DocumentBlock] = []
    headings = ["", "", ""]

    def set_heading(level: int, text: str):
        headings[level] = text[:60]
        for lower in range(level + 1, 3):
            headings[lower] = ""

    def add(item: dict, block_type: str, text: str, rows: list[list[str]], fields: dict[str, str]):
        order = len(blocks) + 1
        blocks.append(DocumentBlock(
            document_id=document_id,
            document_version=version,
            source_block_id=f"B{order:04d}",
            block_type=block_type,
            location=BlockLocation(section=item.get("section", 0), order=order, page=item.get("page"),
                                   heading_path=[h for h in headings if h]),
            text=text,
            rows=rows,
            fields=fields,
        ))

    for item in raw:
        rows = item.get("rows") or []
        if item["type"] == "table":
            if not any(any(c for c in r) for r in rows):
                continue
            # 'Ⅱ | | 제안 요청 내용' 같은 1행짜리 장 제목 표
            flat = [c for c in (rows[0] if len(rows) == 1 else []) if c]
            if flat and len(flat) <= 3 and flat[0] and flat[0][0] in ROMAN:
                set_heading(0, " ".join(flat))
            if _is_requirement_table(rows):
                fields = parse_requirement_table(rows)
                add(item, "requirement_table", render_requirement_table(fields), rows, fields)
            elif columns := _requirement_list_columns(rows):
                # 목록형: 행마다 요구사항 블록 하나 (근거·추출 단위를 요구사항별로 나누기 위해)
                for row in _align_rows(rows[1:], len(rows[0])):
                    fields = {f: row[i].strip() for i, f in columns.items() if row[i].strip()}
                    if REQ_CODE.match(fields.get("code", "")):
                        add(item, "requirement_table", render_requirement_table(fields), [rows[0], row], fields)
            else:
                add(item, "table", render_table(rows), rows, {})
        else:
            text = item["text"]
            level = item.get("heading_level", heading_level(text))
            if level is not None:
                set_heading(level, text)
            add(item, "paragraph", text, rows, {})
    return blocks
