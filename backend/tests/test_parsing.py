"""파싱·청킹 테스트: PDF 표 분리, 인코딩 대체, 미지원 형식 안내, RFP 구조 청커."""
import pytest

from app.agents.tools import chunker
from app.agents.tools.doc_quality import build_sources, check_document
from app.agents.tools.document_parser import UnsupportedFormatError, check_supported, parse_document

RFP_TEXT = """대한적십자사 ITSM 구축 과업지시서

Ⅰ. 사업 개요
1. 사업 목적
□ IT서비스 요청·처리 이력을 통합 관리하기 위한 체계 구축
○ 장애 원인 분석과 선제적 대응

Ⅱ. 과업 내용
1. 기능 요구사항
가. 서비스 요청 관리
□ 사용자는 IT서비스 요청을 등록할 수 있어야 한다.
○ 요청 접수 시 처리 유형과 담당자를 배정한다.
나. 결재 관리
□ 서비스 요청은 부서장 결재 후 처리한다.

Ⅲ. 제안서 작성 및 제출
1. 제출 서류
□ 제안서 10부를 제출한다.
"""


def test_pdf_tables_are_separated_and_requirement_table_recognized(tmp_path):
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    from pdf_fixture import make_rfp_pdf

    blocks = parse_document(make_rfp_pdf(tmp_path / "rfp.pdf"))
    req = [b for b in blocks if b.block_type == "requirement_table"]
    assert len(req) == 1, "병합셀이 있는 '고유번호' 표를 요구사항 표로 인식"
    assert req[0].fields["code"] == "SFR-001" and req[0].fields["acceptance"] == "필수"
    assert req[0].fields["details"] == "사용자는 서비스 요청을 등록한다"
    assert req[0].location.page == 1
    assert req[0].location.heading_path[-1] == "나. 기능 요구사항", "PDF 본문에서도 제목 경로를 잡음"
    long_tables = [b for b in blocks if b.block_type == "table"]
    assert len(long_tables) == 1 and len(long_tables[0].rows) == 61, "페이지를 넘는 표는 하나로 합침"


def test_pdf_decorations_title_box_nested_table_page_number(tmp_path):
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    from pdf_fixture import make_decorated_pdf

    blocks = parse_document(make_decorated_pdf(tmp_path / "decorated.pdf"))
    texts = [b.text for b in blocks]
    assert "Ⅳ 과업내용" in texts, "글자 한 칸씩 + 그림자 장 제목 상자를 최상위 제목 하나로"
    assert blocks[-1].location.heading_path == ["Ⅴ 기타사항"], "표로 잡히지 않은 제목 상자('Ⅴ 기 타 사 항')도 제목"
    assert not any("과과" in t or "업업" in t for t in texts), "그림자로 두 번 찍힌 글자는 하나로"
    req = [b for b in blocks if b.block_type == "requirement_table"]
    assert len(req) == 1 and req[0].location.heading_path == ["Ⅳ 과업내용", "1. 보안 요건"]
    assert "9자 이상" in req[0].fields["details"], "중첩 표 글자는 바깥 표 셀에 포함"
    assert "길이 | 9자 이상" in req[0].fields["details"], "중첩 표는 항목명과 값을 같은 줄에 둔다"
    assert req[0].fields["details"].index("다음 규칙") < req[0].fields["details"].index("길이"), "셀 안 순서 유지"
    assert not [b for b in blocks if b.block_type == "table"], "중첩 표를 별도 표로 다시 꺼내지 않음"
    assert not any(t.strip("- ").isdigit() for t in texts), "꼬리말 쪽번호는 버림"


def test_printed_pdf_page_number_mid_page_and_cut_tables(tmp_path):
    """인쇄한 PDF: 쪽 가운데의 '- 18 -'은 버리고, 그 앞뒤로 잘린 요구사항 표를 하나로 합친다."""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    from pdf_fixture import make_printed_pdf

    blocks = parse_document(make_printed_pdf(tmp_path / "printed.pdf"))
    assert not [b.text for b in blocks if b.block_type == "paragraph" and b.text.strip().startswith("- ")], \
        "쪽 가운데에 찍힌 원래 쪽 번호도 버림"
    req = {b.fields["code"]: b.fields for b in blocks if b.block_type == "requirement_table"}
    assert set(req) == {"SFR-001", "SFR-002", "SFR-003"}
    assert "수정할 수 있어야" in req["SFR-001"]["details"] and req["SFR-001"]["outputs"] == "기능정의서", \
        "빈 첫 칸으로 시작하는 뒷부분"
    assert req["SFR-002"]["name"] == "자동 보고서 생성" and "정기적으로" in req["SFR-002"]["details"] \
        and req["SFR-002"]["outputs"] == "화면설계서", "칸 이름으로 시작하는 뒷부분"
    tables = [b for b in blocks if b.block_type == "table"]
    assert len(tables) == 1 and "알림 방식" in tables[0].text, "'구분 | 내용' 일반 표는 요구사항에 붙이지 않음"
    assert "알림 방식" not in req["SFR-003"]["details"]


def test_pdf_table_without_outer_side_borders(tmp_path):
    """좌우 바깥 세로선이 없는 표도 값 칸을 잡는다 (고객요구사항5.pdf 양식). 위아래 표는 합쳐지지 않는다."""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("reportlab")
    from pdf_fixture import make_open_side_pdf

    blocks = parse_document(make_open_side_pdf(tmp_path / "open.pdf"))
    req = [b for b in blocks if b.block_type == "requirement_table"]
    assert [(b.fields["code"], b.fields["name"]) for b in req] == [("SFR-001", "회원 가입"), ("SER-001", "접근 통제")]
    assert req[0].fields["details"] == "사용자는 회원 가입을 할 수 있다"
    assert req[1].location.heading_path[-1] == "다. 보안 요구사항", "표 사이 제목은 본문으로 남음"


def test_obligations_outside_requirement_tables_are_reported_and_optional(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    path = _make_docx(tmp_path / "rfp.docx")
    doc = Document(path)
    doc.add_heading("2. 일반 요건", level=2)
    for line in ("□ 제안사는 장애 발생 시 2시간 이내에 현장에 도착하여야 함",
                 "□ 제안사는 월 1회 정기점검 결과를 제출하여야 함",
                 "□ 작업 전 보안 서약서를 제출할 것"):
        doc.add_paragraph(line)
    doc.add_heading("Ⅲ. 제안서 작성 요령", level=1)
    doc.add_paragraph("□ 제안서는 10부를 제출하여야 함. 제본하여야 함. 날인할 것")
    doc.save(path)

    blocks = parse_document(path)
    sources = build_sources(blocks)
    assert all(s.code for s in sources), "기본값은 요구사항 표만 추출 단위"
    notes = [f for f in check_document(blocks, sources) if f.kind == "DOC_NARRATIVE_OUTSIDE_TABLES"]
    assert len(notes) == 1 and "2. 일반 요건" in notes[0].message and "작성 요령" not in notes[0].message

    with_narrative = build_sources(blocks, include_narrative=True)
    narrative = [s for s in with_narrative if s.source_key.startswith("N")]
    assert [s.heading for s in narrative] == ["Ⅱ. 과업 내용 > 2. 일반 요건"], "제안서 작성 장은 제외"
    assert not [f for f in check_document(blocks, with_narrative) if f.kind == "DOC_NARRATIVE_OUTSIDE_TABLES"]


def test_text_file_cp949_fallback(tmp_path):
    path = tmp_path / "memo.txt"
    path.write_bytes("Ⅰ. 사업 개요\n\n□ 요청 접수와 담당자 배정 기능을 구축한다.\n".encode("cp949"))
    blocks = parse_document(path)
    assert blocks[0].text == "Ⅰ. 사업 개요"
    assert "담당자 배정" in blocks[1].text


@pytest.mark.parametrize("name,needle", [("old.hwp", "HWPX"), ("slides.pptx", "PDF"), ("bundle.zip", "압축")])
def test_unsupported_format_gives_conversion_hint(tmp_path, name, needle):
    path = tmp_path / name
    path.write_bytes(b"dummy")
    with pytest.raises(UnsupportedFormatError) as caught:
        check_supported(path)
    assert needle in caught.value.hint


def test_rfp_chunker_follows_heading_hierarchy_and_skips_proposal_sections(tmp_path):
    path = tmp_path / "rfp.txt"
    path.write_text(RFP_TEXT, encoding="utf-8")
    blocks = parse_document(path)
    sources = build_sources(blocks)
    headings = [s.heading for s in sources]
    assert "Ⅰ. 사업 개요 > 1. 사업 목적" in headings
    assert "Ⅱ. 과업 내용 > 1. 기능 요구사항 > 가. 서비스 요청 관리" in headings
    assert "Ⅱ. 과업 내용 > 1. 기능 요구사항 > 나. 결재 관리" in headings
    assert not any("제안서" in h for h in headings), "제안서 작성 장은 추출 대상에서 제외"
    kinds = {f.kind: f for f in check_document(blocks, sources)}
    assert "DOC_SECTIONS_SKIPPED" in kinds and "제안서 작성" in kinds["DOC_SECTIONS_SKIPPED"].message


def test_rfp_chunker_splits_oversize_section_at_item_marks(tmp_path, monkeypatch):
    monkeypatch.setattr(chunker, "MAX_CHARS", 120)
    items = "\n".join(f"□ 항목 {i} 요구사항 설명입니다 " + "가" * 30 for i in range(1, 7))
    path = tmp_path / "long.txt"
    path.write_text(f"Ⅱ. 과업 내용\n1. 기능 요구사항\n{items}\n", encoding="utf-8")
    blocks = parse_document(path)
    plan = chunker.plan_chunks(blocks)
    assert len(plan.sources) > 1, "상한을 넘는 절은 나눈다"
    by_id = {b.source_block_id: b for b in blocks}
    for src in plan.sources[1:]:
        assert by_id[src.block_ids[0]].text.startswith("□"), "나눈 지점은 □ 항목 시작"
    assert all(s.heading == "Ⅱ. 과업 내용 > 1. 기능 요구사항" for s in plan.sources)


def test_kobaco_rfp_requirement_tables_and_structure(kobaco_path):
    """요구사항 표 형식이 다른 RFP: '요구사항 ID | MAR-001 | 요구사항 명칭 | …', 목차 쪽번호 줄, [붙임]·[별지] 서식."""
    if not kobaco_path.exists():
        pytest.skip("KOBACO RFP 없음")
    from app.agents.tools.doc_quality import prefix_categories_from_summary

    blocks = parse_document(kobaco_path)
    req = [b for b in blocks if b.block_type == "requirement_table"]
    assert len(req) == 80
    first = req[0].fields
    assert (first["code"], first["name"]) == ("MAR-001", "업무시스템 유지관리")
    assert {"category", "definition", "details", "outputs"} <= set(first)
    toc = next(b for b in blocks if b.text.startswith("1. 사업 개요") and b.text.endswith("4"))
    assert toc.location.heading_path == [], "목차 줄(끝에 쪽번호)은 제목이 아님"
    assert req[0].location.heading_path == ["4. 제안요청 내용", "다. 요구사항별 상세 내용"]
    assert prefix_categories_from_summary(blocks)["MAR"] == "유지관리 수행 요구사항"

    sources = build_sources(blocks)
    kinds = [f.kind for f in check_document(blocks, sources)]
    assert len(sources) == 80 and "DOC_HEADING_CATEGORY_MISMATCH" not in kinds
    assert "DOC_NAME_MISMATCH" in kinds, "MAR-035 목록·상세 명칭 차이"
    plan = chunker.plan_chunks([b for b in blocks if b.block_type != "requirement_table"])
    assert not any("별지" in s.heading or "붙임" in s.heading for s in plan.sources), "서식·부록은 건너뜀"


def _make_docx(path):
    from docx import Document

    doc = Document()
    doc.add_heading("Ⅱ. 과업 내용", level=1)
    doc.add_heading("1. 기능 요구사항", level=2)
    kv = doc.add_table(rows=3, cols=4)
    for cell, text in zip(kv.rows[0].cells, ["요구사항 ID", "REQ-001", "요구사항 명칭", "서비스 요청 등록"]):
        cell.text = text
    kv.cell(1, 0).text, kv.cell(2, 0).text = "요구사항 분류", "세부내용"
    kv.cell(1, 1).merge(kv.cell(1, 3)).text = "기능 요구사항"
    kv.cell(2, 1).merge(kv.cell(2, 3)).text = "사용자는 IT서비스 요청을 등록한다."
    doc.add_paragraph("□ 요청 처리 요구사항 목록")
    lst = doc.add_table(rows=3, cols=4)
    for cell, text in zip(lst.rows[0].cells, ["구분", "요구사항ID", "요구사항명", "세부내용"]):
        cell.text = text
    for r, row in enumerate([["REQ-101", "요청 접수", "서비스데스크 담당자는 등록된 요청을 접수하거나 반려한다."],
                             ["REQ-102", "담당자 배정", "접수된 요청에 처리 유형과 1·2·3순위 담당자를 배정한다."]], start=1):
        for c, text in enumerate(row, start=1):
            lst.cell(r, c).text = text
    lst.cell(1, 0).merge(lst.cell(2, 0)).text = "기능"     # 세로 병합된 분류 칸
    doc.save(path)
    return path


def test_docx_headings_requirement_tables_and_list_rows(tmp_path):
    pytest.importorskip("docx")
    blocks = parse_document(_make_docx(tmp_path / "rfp.docx"))
    req = [b for b in blocks if b.block_type == "requirement_table"]
    assert [b.fields["code"] for b in req] == ["REQ-001", "REQ-101", "REQ-102"], "상세형 1개 + 목록형 2행"
    assert req[0].fields["category"] == "기능 요구사항" and "등록" in req[0].fields["details"]
    assert req[2].fields["category"] == "기능", "세로 병합된 분류 칸 값을 이어받음"
    assert req[2].fields["name"] == "담당자 배정"
    assert req[0].location.heading_path == ["Ⅱ. 과업 내용", "1. 기능 요구사항"], "Word 제목 스타일로 제목 경로"
    assert [s.source_key for s in build_sources(blocks)] == ["REQ-001", "REQ-101", "REQ-102"]


def test_parse_command_on_folder(tmp_path, capsys):
    pytest.importorskip("docx")
    from app.cli import main

    _make_docx(tmp_path / "a.docx")
    (tmp_path / "b.txt").write_text("Ⅰ. 사업 개요\n□ 요청 관리 시스템을 구축한다.\n", encoding="utf-8")
    (tmp_path / "c.hwp").write_bytes(b"x")
    assert main(["parse", str(tmp_path), "--output-dir", str(tmp_path / "out")]) == 0
    out = capsys.readouterr().out
    assert "[완료] a.docx" in out and "[완료] b.txt" in out
    assert "[건너뜀] c.hwp" in out and "HWPX" in out
